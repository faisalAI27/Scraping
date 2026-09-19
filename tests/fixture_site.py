"""Deterministic test site; access is granted to one loopback origin only."""

from collections import Counter
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import io
import json
import threading
from urllib.parse import urlsplit

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

HOME = b"""<!doctype html><html><head><title>Harbor Services</title></head><body>
<nav><a href="/service">Services</a><a href="/service?utm_source=nav">Services again</a></nav>
<main><h1>Harbor Services</h1><p>Returns are not accepted after 30 days, except faulty items.</p>
<p>Plan A costs $19.95 per month.</p><p>Plan B costs $29.95 per month.</p>
<div class="service-card">Repairs available by appointment.</div>
<h2>Specifications</h2><table><caption>Model specifications</caption>
<thead><tr><th>Model</th><th>Capacity (kg)</th><th>Price</th></tr></thead>
<tr><td>AB-100-X</td><td>25 kg</td><td>$199.00</td></tr></table>
<h2>Questions</h2><details><summary>Are Sundays open?</summary><p>No, except emergency repairs.</p></details>
<ol start="3"><li>Switch off power.</li><li>Do not remove the seal.</li></ol>
<a href="/brochure.pdf">Brochure</a><a href="/policy.docx">Policy</a>
<a href="/scanned.pdf">Scanned notice</a><a href="/mixed.pdf">Mixed PDF</a>
<img src="/notice.png" alt="Public service hours notice" width="1100" height="220">
<img src="/logo.png" alt="Logo" width="20" height="20">
<a href="/js">Live availability</a><a href="/short">Short page</a><a href="/transient">Transient</a>
<a href="/failure">Failure</a><a href="/large">Large</a><a href="/denied">Robots denied</a>
<a href="/redirect-private">Bad redirect</a><a href="/rate">Rate limited</a>
<a href="/old.doc">Old format</a><a href="https://outside.invalid/">Outside</a>
<a href="/?page=2">Next page</a><a href="#questions">Same page</a></main>
<div class="cookie-banner">Accept cookies to continue</div>
<footer><address>Contact: info@harbor.example<br>Monday-Friday 09:00-17:00</address></footer></body></html>"""

FACTS = [
    "Returns are not accepted after 30 days, except faulty items.",
    "Plan A costs $19.95 per month.",
    "Plan B costs $29.95 per month.",
    "AB-100-X",
    "25 kg",
    "$199.00",
    "No, except emergency repairs.",
    "Do not remove the seal.",
    "Monday-Friday 09:00-17:00",
    "info@harbor.example",
    "Warranty lasts 24 months.",
    "Cancellation is not allowed except during the first 7 days.",
    "Public notice: Open Saturday 10:00-14:00.",
    "Appointments available Tuesday.",
]


def binary_fixtures():
    image = Image.new("RGB", (1500, 260), "white")
    draw = ImageDraw.Draw(image)
    font = (
        ImageFont.truetype("Arial.ttf", 44)
        if __import__("sys").platform == "darwin"
        else ImageFont.truetype("DejaVuSans.ttf", 44)
    )
    draw.text((35, 95), "Public notice: Open Saturday 10:00-14:00.", font=font, fill="black")
    image_data = io.BytesIO()
    image.save(image_data, format="PNG")
    brochure = io.BytesIO()
    pdf = canvas.Canvas(brochure, pagesize=(612, 792))
    pdf.setTitle("Service brochure")
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(50, 735, "Service brochure")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(50, 700, "Warranty lasts 24 months.")
    pdf.drawString(50, 680, "Model AB-100-X capacity: 25 kg.")
    pdf.save()
    scanned = io.BytesIO()
    pdf = canvas.Canvas(scanned, pagesize=(800, 300))
    pdf.drawImage(ImageReader(image), 25, 80, width=750, height=130)
    pdf.save()
    mixed = io.BytesIO()
    pdf = canvas.Canvas(mixed, pagesize=(800, 400))
    pdf.setFont("Helvetica", 14)
    pdf.drawString(25, 350, "Mixed brochure: Service is not available overnight.")
    pdf.drawImage(ImageReader(image), 25, 80, width=750, height=130)
    pdf.save()
    document = Document()
    document.core_properties.title = "Cancellation policy"
    document.add_heading("Cancellation policy", 0)
    document.add_paragraph("Cancellation is not allowed except during the first 7 days.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Service", "Fee"
    table.cell(1, 0).text, table.cell(1, 1).text = "Installation", "$49.00"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph("Prices do not include optional delivery.")
    docx = io.BytesIO()
    document.save(docx)
    return {
        "/brochure.pdf": (brochure.getvalue(), "application/pdf"),
        "/scanned.pdf": (scanned.getvalue(), "application/pdf"),
        "/mixed.pdf": (mixed.getvalue(), "application/pdf"),
        "/policy.docx": (
            docx.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        "/notice.png": (image_data.getvalue(), "image/png"),
    }


@contextmanager
def fixture_site():
    counts = Counter()
    fixtures = binary_fixtures()

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_):
            pass

        def do_GET(self):
            counts[self.path] += 1
            path = urlsplit(self.path).path
            status, body, mime, headers = 200, b"", "text/html", {}
            if path == "/robots.txt":
                body, mime = b"User-agent: *\nDisallow: /denied\n", "text/plain"
            elif path == "/sitemap.xml":
                body, mime = (
                    f"<sitemapindex><sitemap><loc>{site_origin}/child.xml</loc></sitemap></sitemapindex>".encode(),
                    "application/xml",
                )
            elif path == "/child.xml":
                body, mime = (
                    f"<urlset><url><loc>{site_origin}/service</loc></url><url><loc>{site_origin}/?page=2</loc></url></urlset>".encode(),
                    "application/xml",
                )
            elif self.path == "/?page=2":
                body = b"<h1>Second page</h1><p>Policy version 2: Returns are accepted after 30 days.</p>"
            elif path == "/":
                body = HOME
            elif path == "/service":
                body = b"<main><h1>Services</h1><p>Installation includes testing.</p></main>"
            elif path == "/short":
                body = b"<p>Closed Sundays.</p>"
            elif path == "/js-chain":
                body = b"""<main id="app">Loading...</main><script>
                fetch('/api/message').then(r=>r.json()).then(()=>fetch('/short'))
                .then(r=>r.text()).then(text=>document.getElementById('app').innerHTML=text);
                </script>"""
            elif path == "/challenge":
                body = b"<title>Just a moment...</title><script>window._cf_chl_opt = {};</script><p>Enable JavaScript and cookies to continue.</p>"
            elif path == "/js":
                body = b"""<html><body><main id="app">Loading...</main><script>
                fetch('/api/message').then(r=>r.json()).then(d=>document.getElementById('app').innerHTML='<h1>Availability</h1><p>'+d.text+'</p><a href="/rendered-only">Details</a>');
                fetch('http://127.0.0.1:1/private');
                fetch('/post-target', {method:'POST',body:'should be blocked'});
                </script></body></html>"""
            elif path == "/api/message":
                body, mime = (
                    json.dumps({"text": "Appointments available Tuesday."}).encode(),
                    "application/json",
                )
            elif path == "/rendered-only":
                body = b"<p>Rendered link was discovered.</p>"
            elif path == "/transient":
                status = 503 if counts[self.path] == 1 else 200
                body = b"<p>Recovered transient service.</p>"
            elif path == "/rate":
                status = 429 if counts[self.path] == 1 else 200
                headers["Retry-After"] = "0"
                body = b"<p>Recovered after rate limit.</p>"
            elif path == "/failure":
                status, body = 503, b"Unavailable"
            elif path == "/large":
                body = b"x" * 120000
            elif path == "/redirect-private":
                status, headers["Location"] = 302, "http://127.0.0.1:1/private"
            elif path == "/old.doc":
                body, mime = b"legacy-document", "application/msword"
            elif path in fixtures:
                body, mime = fixtures[path]
            else:
                status, body = 404, b"Not found"
            self.send_response(status)
            self.send_header("Content-Type", mime)
            self.send_header("Content-Length", str(len(body)))
            for k, v in headers.items():
                self.send_header(k, v)
            self.end_headers()
            try:
                self.wfile.write(body)
            except (BrokenPipeError, ConnectionResetError):
                pass

        def do_POST(self):
            counts["POST " + self.path] += 1
            self.send_response(400)
            self.end_headers()

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    site_origin = f"http://127.0.0.1:{server.server_port}"
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield site_origin, counts, fixtures
    finally:
        server.shutdown()
        server.server_close()
        thread.join()
